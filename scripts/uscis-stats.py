"""Builds src/content/tools/uscis-stats.json from the official USCIS report
"Number of Service-wide Forms by Quarter, Form Status, and Processing Time".

USCIS publishes one spreadsheet per quarter at
https://www.uscis.gov/tools/reports-and-studies/immigration-and-citizenship-data
(file names such as quarterly_all_forms_fy2026_q3_v1.xlsx). Download the
quarters you want into one folder and run:

    python scripts/uscis-stats.py <folder>

When a quarter's file is not published, its received, approved, denied and
completed counts are derived from the fiscal-year-to-date columns of the next
report in the same fiscal year. Pending counts and median times cannot be
derived, so they stay empty for that quarter. Needs openpyxl.
"""

import glob
import json
import os
import re
import sys

import openpyxl

# Rows the website uses: id, short label, form, group, and a pattern for the
# row's form column and title column. Titles changed slightly between reports.
ROWS = [
    ("i130-ir", "I-130, immediate relatives", "I-130", "Family", r"^I-130", r"Alien Relative \(Immediate"),
    ("i130-other", "I-130, other relatives", "I-130", "Family", r"^I-130", r"Alien Relative \(All Other"),
    ("i129f", "I-129F, fiancé(e)", "I-129F", "Family", r"^I-129F", r"Fianc"),
    ("i601a", "I-601A, provisional waiver", "I-601A", "Family", r"^I-601A", r"Provisional"),
    ("i751", "I-751, remove conditions", "I-751", "Family", r"^I-751", r"Remove Conditions"),
    ("i129", "I-129, nonimmigrant worker", "I-129", "Employment", r"^I-129(?!F)", r"Nonimmigrant Worker"),
    ("i140", "I-140, immigrant worker", "I-140", "Employment", r"^I-140", r"Alien Worker"),
    ("i526e", "I-526E, regional center investor", "I-526E", "Employment", r"^I-526E", r"Regional Center Investor"),
    ("i829", "I-829, investor conditions", "I-829", "Employment", r"^I-829", r"Investor to Remove"),
    ("i765-aos", "I-765, pending green card", "I-765", "Work and travel", r"^I-765", r"\(Adjustment"),
    ("i765-asylum", "I-765, asylum", "I-765", "Work and travel", r"^I-765", r"\(Asylum\)"),
    ("i765-other", "I-765, all other", "I-765", "Work and travel", r"^I-765", r"\(All Other\)"),
    ("i131-ap", "I-131, advance parole", "I-131", "Work and travel", r"^I-131", r"Advance Parole"),
    ("i131-td", "I-131, travel documents", "I-131", "Work and travel", r"^I-131$", r"^Application for Travel Document(s)?(, Parole Documents, and Arrival/Departure Records)?$"),
    ("i539", "I-539, extend or change status", "I-539", "Work and travel", r"^I-539", r"Extend/Change"),
    ("i589", "I-589, affirmative asylum", "I-589", "Humanitarian", r"^I-589", r"Asylum and for Withholding"),
    ("i730", "I-730, refugee or asylee relative", "I-730", "Humanitarian", r"^I-730", r"Refugee/Asylee Relative"),
    ("i821", "I-821, Temporary Protected Status", "I-821", "Humanitarian", r"^I-821(?!D)", r"Temporary Protected"),
    ("i914", "I-914, T visa", "I-914", "Humanitarian", r"^I-914", r"T Nonimmigrant"),
    ("i918", "I-918, U visa", "I-918", "Humanitarian", r"^I-918", r"U Nonimmigrant"),
    ("i360", "I-360, special immigrant and VAWA", "I-360", "Humanitarian", r"^I-360", r"Special Immigrant"),
    ("i485-family", "I-485, family", "I-485", "Green card", r"^I-485", r"\(Family\)"),
    ("i485-employment", "I-485, employment", "I-485", "Green card", r"^I-485", r"\(Employment\)"),
    ("i485-asylum", "I-485, asylee", "I-485", "Green card", r"^I-485", r"\(Asylum\)"),
    ("i485-refugee", "I-485, refugee", "I-485", "Green card", r"^I-485", r"\(Refugee\)"),
    ("i90", "I-90, card renewal", "I-90", "Green card", r"^I-90$", r"Replace Permanent Resident Card"),
    ("n400", "N-400, naturalization", "N-400", "Citizenship", r"^N-400", r"^Application for Naturalization$"),
    ("n600", "N-600, certificate of citizenship", "N-600", "Citizenship", r"^N-600", r"Certificate of Citizenship"),
    ("i290b", "I-290B, appeal or motion", "I-290B", "Appeals and waivers", r"^I-290B", r"Appeal or Motion"),
    ("waivers", "Waivers (I-601, I-212, I-612 and others)", "Waivers", "Appeals and waivers", r"^Waivers", r"^Waivers$"),
]

FIELDS = ["received", "approved", "denied", "completed", "pending", "median"]


def num(v):
    if isinstance(v, (int, float)):
        return v
    return None


def read(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    period = next(str(r[0]) for r in rows[:6] if r and r[0] and re.search(r"\d{4}", str(r[0])))
    out = {}
    for r in rows:
        if not r or not r[0] or not r[1]:
            continue
        # Footnote numbers are glued to the form code (I-58912 is I-589, note 12).
        code, title = str(r[0]).strip(), str(r[1]).strip()
        for rid, _, form, _, _, title_re in ROWS:
            if rid in out or not code.startswith(form) or not re.search(title_re, title):
                continue
            q = dict(zip(FIELDS, [num(x) for x in r[2:8]]))
            # USCIS prints 0 where it reports no median time, as for affirmative asylum.
            q["median"] = round(q["median"], 1) if q["median"] else None
            ytd = dict(zip(FIELDS[:5], [num(x) for x in r[8:13]]))
            out[rid] = {"q": q, "ytd": ytd}
    total = next((r for r in rows if r and str(r[0]).strip().upper() == "TOTAL"), None)
    if total:
        out["total"] = {"q": dict(zip(FIELDS, [num(x) for x in total[2:8]])), "ytd": dict(zip(FIELDS[:5], [num(x) for x in total[8:13]]))}
    return period, out


def main(folder):
    reports = {}
    for f in glob.glob(os.path.join(folder, "quarterly_all_forms_fy*_q*.xlsx")):
        m = re.search(r"fy(\d{4})_q(\d)", os.path.basename(f))
        reports[(int(m.group(1)), int(m.group(2)))] = (os.path.basename(f), *read(f))
    if not reports:
        sys.exit("No quarterly_all_forms files found")

    keys = sorted(reports)
    first, last = keys[0], keys[-1]
    quarters = []
    fy, q = first
    while (fy, q) <= last:
        quarters.append((fy, q))
        fy, q = (fy, q + 1) if q < 4 else (fy + 1, 1)

    ids = [r[0] for r in ROWS] + ["total"]
    series = {i: [] for i in ids}
    meta = []
    for fy, q in quarters:
        if (fy, q) in reports:
            name, period, data = reports[(fy, q)]
            meta.append({"fy": fy, "q": q, "label": f"FY{fy % 100} Q{q}", "period": period, "source": name, "derived": False})
            for i in ids:
                series[i].append(data.get(i, {}).get("q"))
            continue
        # Missing report: derive flows from the next report's year-to-date totals.
        later = [k for k in keys if k[0] == fy and k[1] > q]
        if not later:
            meta.append({"fy": fy, "q": q, "label": f"FY{fy % 100} Q{q}", "period": None, "source": None, "derived": True})
            for i in ids:
                series[i].append(None)
            continue
        nxt = later[0]
        name, _, data = reports[nxt]
        meta.append({"fy": fy, "q": q, "label": f"FY{fy % 100} Q{q}", "period": None, "source": name, "derived": True})
        for i in ids:
            row = data.get(i)
            if not row:
                series[i].append(None)
                continue
            vals = {}
            for f in FIELDS[:4]:
                total = row["ytd"].get(f)
                parts = [row["q"].get(f)]
                for k in range(1, nxt[1]):
                    if k == q:
                        continue
                    other = reports.get((fy, k))
                    parts.append(other[2].get(i, {}).get("q", {}).get(f) if other else None)
                vals[f] = None if total is None or any(p is None for p in parts) else max(0, total - sum(parts))
            vals["pending"] = None
            vals["median"] = None
            series[i].append(vals)

    out = {
        "title": "Number of Service-wide Forms by Quarter, Form Status, and Processing Time",
        "publisher": "U.S. Citizenship and Immigration Services, Office of Performance and Quality",
        "url": "https://www.uscis.gov/tools/reports-and-studies/immigration-and-citizenship-data",
        "quarters": meta,
        "rows": [
            {"id": rid, "label": label, "form": form, "group": group, "series": series[rid]}
            for rid, label, form, group, _, _ in ROWS
        ],
        "total": series["total"],
    }
    dest = os.path.join(os.path.dirname(__file__), "..", "src", "content", "tools", "uscis-stats.json")
    with open(dest, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, separators=(",", ":"))
    for rid, label, *_ in ROWS:
        cov = "".join("." if s is None else ("d" if m["derived"] else "#") for s, m in zip(series[rid], meta))
        print(f"{rid:16} {cov}")
    print("wrote", os.path.normpath(dest))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else ".")
